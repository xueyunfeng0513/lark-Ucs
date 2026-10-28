from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
import pytest
import xlwt
from openpyxl.styles import PatternFill

from excel_counter.excel import WorkbookError, count_rows


def test_xlsx_counts_content_not_formatting(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["来源订单号", "数量"])
    sheet.append(["A001", 0])
    sheet.append([None, None])
    sheet.append([" ", "\t"])
    sheet.append([False, None])
    sheet.append(["=1+1", None])
    sheet.row_dimensions[2].hidden = True
    sheet["A10000"].fill = PatternFill("solid", fgColor="FF0000")
    book.create_sheet("空表")
    hidden = book.create_sheet("隐藏表")
    hidden.sheet_state = "hidden"
    hidden.append(["hidden"])
    path = tmp_path / "orders.xlsx"
    book.save(path)
    assert count_rows(path) == {"订单": 4, "空表": 0, "隐藏表": 1}


def test_xls_counts_all_sheets_and_zero(tmp_path):
    book = xlwt.Workbook()
    sheet = book.add_sheet("旧版表")
    sheet.write(0, 0, "表头")
    sheet.write(1, 0, 0)
    sheet.write(2, 0, " ")
    sheet.write(4, 0, False)
    book.add_sheet("空表")
    path = tmp_path / "orders.xls"
    book.save(str(path))
    assert count_rows(path) == {"旧版表": 3, "空表": 0}


def test_2000_rows_and_header(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["编号", "数量"])
    for i in range(2000):
        sheet.append([f"A{i:04}", i])
    path = tmp_path / "2000.xlsx"
    book.save(path)
    assert count_rows(path) == {"Sheet": 2001}


@pytest.mark.parametrize("suffix", [".xlsx", ".xls", ".xlsm"])
def test_invalid_file_is_safe_error(tmp_path, suffix):
    path = tmp_path / ("bad" + suffix)
    path.write_text("not a workbook")
    with pytest.raises(WorkbookError):
        count_rows(path)


def test_zip_expansion_limit(tmp_path):
    path = tmp_path / "bomb.xlsx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "x" * 10000)
    with pytest.raises(WorkbookError, match="解压"):
        count_rows(path, max_expanded_bytes=100)


def test_sparse_rows_have_scan_limit(tmp_path):
    book = openpyxl.Workbook()
    book.active["A100"] = "far away"
    path = tmp_path / "sparse.xlsx"
    book.save(path)
    with pytest.raises(WorkbookError, match="扫描"):
        count_rows(path, max_scan_rows=10)


def test_stale_dimension_does_not_hide_data(tmp_path):
    book = openpyxl.Workbook()
    book.active.append(["header"])
    book.active.append(["record"])
    path = tmp_path / "source.xlsx"
    book.save(path)
    bad_dimension = tmp_path / "bad-dimension.xlsx"
    with ZipFile(path) as original, ZipFile(bad_dimension, "w") as fixed:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:A2"', b'ref="A1:A1"')
            fixed.writestr(item, content)
    assert count_rows(bad_dimension) == {"Sheet": 2}
